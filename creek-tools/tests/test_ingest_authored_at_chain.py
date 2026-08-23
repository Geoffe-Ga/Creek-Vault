"""The shared FEAT-031 authored-at extraction chain (#855, #856, #870).

Every format-specific ingestor implements the same four-step shape:

1. open a format-specific metadata source — the optional dependency may
   be absent and the file may be malformed;
2. walk an **ordered** candidate chain (``created`` before ``modified``,
   ``creationdate`` before ``moddate``, ``\\creatim`` before ``\\revtim``);
3. swallow the parse failure and fall through to the next candidate;
4. return ``None`` rather than guess.

Those are exactly the arms that never get exercised by happy-path
fixtures, and they are the arms that decide whether a real date is
silently lost. This module drives the three shared failure conditions
— *optional library absent*, *metadata absent*, *metadata malformed* —
across the formats that share the pattern, plus the ordering assertions
that a reordering mutant must not survive.
"""

from __future__ import annotations

import sys
from datetime import UTC, datetime, timedelta, timezone
from typing import TYPE_CHECKING, Any

import pytest

from creek.ingest import documents, presentations, spreadsheets
from creek.ingest.base import safe_parse_authored_at
from creek.ingest.documents import (
    _parse_docx_authored_at,
    _parse_pdf_authored_at,
    _parse_pdf_date,
    _parse_rtf_authored_at,
)
from creek.ingest.presentations import (
    PresentationIngestor,
    PythonPptxBackend,
    _extract_pptx_authored_at,
    _open_pptx_core_properties,
)
from creek.ingest.spreadsheets import (
    OpenpyxlBackend,
    SpreadsheetIngestor,
    _extract_xlsx_authored_at,
    _open_xlsx,
)

if TYPE_CHECKING:
    from pathlib import Path


# ---- The shared helper --------------------------------------------------


class TestSafeParseAuthoredAt:
    """``safe_parse_authored_at`` is the single swallow point."""

    @pytest.mark.parametrize(
        "candidate",
        [
            pytest.param(None, id="none"),
            pytest.param("", id="empty-string"),
            pytest.param("   ", id="whitespace-only"),
        ],
    )
    def test_absent_candidate_is_none(self, candidate: object) -> None:
        """An absent value is not an error — it is the next-candidate signal."""
        assert safe_parse_authored_at(candidate) is None

    @pytest.mark.parametrize(
        "candidate",
        [
            pytest.param("not-a-date", id="prose"),
            pytest.param("2024-13-45", id="impossible-date"),
            pytest.param(object(), id="unparseable-object"),
        ],
    )
    def test_unparseable_candidate_is_swallowed(self, candidate: object) -> None:
        """A ``ValueError`` from the parser must not escape to the caller.

        This is the arm that dies if the ``except ValueError`` is
        removed: the parser raises, and without the swallow the whole
        ingest of that file fails instead of falling through.
        """
        assert safe_parse_authored_at(candidate) is None

    def test_aware_datetime_passes_through_unchanged(self) -> None:
        """An already-aware datetime keeps its own offset."""
        tokyo = timezone(timedelta(hours=9))
        value = datetime(2024, 3, 15, 8, 30, tzinfo=tokyo)
        result = safe_parse_authored_at(value)
        assert result is not None
        assert result == value
        assert result.utcoffset() == timedelta(hours=9)

    def test_naive_datetime_is_localised_to_utc(self) -> None:
        """FEAT-031 defaults a naive source date to UTC, never to local."""
        result = safe_parse_authored_at(datetime(2024, 3, 15, 8, 30))
        assert result == datetime(2024, 3, 15, 8, 30, tzinfo=UTC)

    def test_iso_string_is_parsed(self) -> None:
        """The happy path still resolves, so the guard is not over-broad."""
        result = safe_parse_authored_at("2024-03-15T08:30:00+00:00")
        assert result == datetime(2024, 3, 15, 8, 30, tzinfo=UTC)

    def test_the_two_ingestors_share_one_implementation(self) -> None:
        """Spreadsheets and presentations must not re-grow private copies.

        #855 deleted two byte-identical private helpers in favour of
        this one. A future edit that reinstates either copy — the way
        the duplication arose the first time — fails here.
        """
        assert spreadsheets.safe_parse_authored_at is safe_parse_authored_at
        assert presentations.safe_parse_authored_at is safe_parse_authored_at
        assert not hasattr(spreadsheets, "_safe_parse_authored_at")
        assert not hasattr(presentations, "_safe_parse_authored_at")


# ---- Condition 1: the optional library is absent ------------------------


class TestOptionalDependencyAbsent:
    """A missing optional dep yields ``None``, never a crash or a guess."""

    def test_openpyxl_absent_yields_no_workbook(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """``_open_xlsx`` returns ``None`` when ``openpyxl`` cannot import.

        ``monkeypatch.setitem(sys.modules, ..., None)`` makes the import
        statement itself raise ``ImportError``, which is the only way to
        reach this arm — patching the module attribute would bypass it.
        """
        monkeypatch.setitem(sys.modules, "openpyxl", None)
        assert _open_xlsx(tmp_path / "book.xlsx") is None

    def test_openpyxl_absent_yields_no_authored_at(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """The whole XLSX chain degrades to ``None``, not to mtime."""
        xlsx = tmp_path / "book.xlsx"
        xlsx.write_bytes(b"not really a workbook")
        monkeypatch.setitem(sys.modules, "openpyxl", None)
        assert _extract_xlsx_authored_at(xlsx) is None

    def test_openpyxl_absent_reports_backend_unavailable(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """``is_available()`` is ``False`` exactly when the import fails."""
        monkeypatch.setitem(sys.modules, "openpyxl", None)
        assert OpenpyxlBackend().is_available() is False

    def test_openpyxl_present_reports_backend_available(self) -> None:
        """The positive arm: a real environment reports ``True``."""
        assert OpenpyxlBackend().is_available() is True

    def test_python_pptx_absent_yields_no_core_properties(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """``_open_pptx_core_properties`` returns ``None`` without python-pptx."""
        monkeypatch.setitem(sys.modules, "pptx", None)
        assert _open_pptx_core_properties(tmp_path / "deck.pptx") is None

    def test_python_pptx_absent_yields_no_authored_at(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """The whole PPTX chain degrades to ``None``."""
        deck = tmp_path / "deck.pptx"
        deck.write_bytes(b"not really a deck")
        monkeypatch.setitem(sys.modules, "pptx", None)
        assert _extract_pptx_authored_at(deck) is None

    def test_python_pptx_absent_reports_backend_unavailable(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """``is_available()`` is ``False`` exactly when the import fails."""
        monkeypatch.setitem(sys.modules, "pptx", None)
        assert PythonPptxBackend().is_available() is False

    def test_python_pptx_present_reports_backend_available(self) -> None:
        """The positive arm: a real environment reports ``True``."""
        assert PythonPptxBackend().is_available() is True


# ---- Condition 2: the file is malformed ---------------------------------


class TestMalformedSource:
    """A corrupt file loses its date, not the caller's process."""

    def test_corrupt_xlsx_yields_none(self, tmp_path: Path) -> None:
        """openpyxl raises a wide error set on junk; all of it is swallowed."""
        xlsx = tmp_path / "corrupt.xlsx"
        xlsx.write_bytes(b"PK\x03\x04 definitely not a workbook")
        assert _extract_xlsx_authored_at(xlsx) is None

    def test_missing_xlsx_yields_none(self, tmp_path: Path) -> None:
        """A path that does not exist is a file-open failure, not a raise."""
        assert _extract_xlsx_authored_at(tmp_path / "gone.xlsx") is None

    def test_non_xlsx_suffix_short_circuits(self, tmp_path: Path) -> None:
        """CSV carries no core properties, so the chain never opens it."""
        csv_path = tmp_path / "rows.csv"
        csv_path.write_text("a,b\n1,2\n", encoding="utf-8")
        assert _extract_xlsx_authored_at(csv_path) is None

    def test_corrupt_pptx_yields_none(self, tmp_path: Path) -> None:
        """python-pptx raises a wide error set on junk; all of it is swallowed."""
        deck = tmp_path / "corrupt.pptx"
        deck.write_bytes(b"PK\x03\x04 definitely not a deck")
        assert _open_pptx_core_properties(deck) is None

    def test_core_properties_attribute_error_yields_none(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A python-pptx object without ``core_properties`` degrades cleanly."""

        class _NoProps:
            """Presentation stand-in whose ``core_properties`` is missing."""

            @property
            def core_properties(self) -> Any:
                """Raise, as an older python-pptx build would."""
                raise AttributeError("core_properties")

        import pptx

        monkeypatch.setattr(pptx, "Presentation", lambda _path: _NoProps())
        assert _open_pptx_core_properties(tmp_path / "deck.pptx") is None

    def test_presentation_title_attribute_error_yields_none(self) -> None:
        """``_extract_title`` coerces a raising core-properties to ``None``."""

        class _Raising:
            """Core-properties stand-in whose ``title`` raises."""

            @property
            def title(self) -> str:
                """Raise, as a malformed deck's properties would."""
                raise ValueError("bad title")

        class _Prs:
            """Presentation stand-in exposing the raising properties."""

            core_properties = _Raising()

        assert PythonPptxBackend._extract_title(_Prs()) is None

    @pytest.mark.parametrize(
        "title",
        [
            pytest.param(None, id="none"),
            pytest.param("", id="empty"),
            pytest.param("   ", id="whitespace-only"),
        ],
    )
    def test_blank_presentation_title_is_none(self, title: str | None) -> None:
        """Blank titles collapse to ``None`` so callers can fall back."""

        class _Props:
            """Core-properties stand-in holding one title value."""

            def __init__(self, value: str | None) -> None:
                """Store the title verbatim, blank or not."""
                self.title = value

        class _Prs:
            """Presentation stand-in exposing the core properties."""

            def __init__(self, value: str | None) -> None:
                """Wire up the core properties around *value*."""
                self.core_properties = _Props(value)

        assert PythonPptxBackend._extract_title(_Prs(title)) is None


# ---- Condition 3: the ordered candidate chain ---------------------------


class _FakeProps:
    """Core-properties stand-in with settable ``created`` / ``modified``."""

    def __init__(self, created: object, modified: object) -> None:
        """Store the two candidate values verbatim."""
        self.created = created
        self.modified = modified


class _FakeWorkbook:
    """Workbook stand-in that records whether it was closed."""

    def __init__(self, created: object, modified: object) -> None:
        """Expose ``properties`` and start out un-closed."""
        self.properties = _FakeProps(created, modified)
        self.closed = False

    def close(self) -> None:
        """Record the close so the ``finally`` arm can be asserted."""
        self.closed = True


class TestXlsxCandidateChain:
    """``created`` outranks ``modified``, and the handle is always closed."""

    def test_created_wins_over_modified(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """The chain is ordered; ``modified`` is only a fallback."""
        book = _FakeWorkbook(
            created=datetime(2024, 3, 15, tzinfo=UTC),
            modified=datetime(2020, 1, 1, tzinfo=UTC),
        )
        monkeypatch.setattr(spreadsheets, "_open_xlsx", lambda _path: book)
        result = _extract_xlsx_authored_at(tmp_path / "book.xlsx")
        assert result == datetime(2024, 3, 15, tzinfo=UTC)
        assert book.closed is True

    def test_falls_through_to_modified_when_created_absent(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """An absent ``created`` hands off to ``modified``."""
        book = _FakeWorkbook(created=None, modified=datetime(2020, 1, 1, tzinfo=UTC))
        monkeypatch.setattr(spreadsheets, "_open_xlsx", lambda _path: book)
        assert _extract_xlsx_authored_at(tmp_path / "book.xlsx") == datetime(
            2020, 1, 1, tzinfo=UTC
        )

    def test_falls_through_when_created_is_unparseable(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A junk ``created`` is swallowed, not propagated."""
        book = _FakeWorkbook(
            created="not-a-date", modified=datetime(2020, 1, 1, tzinfo=UTC)
        )
        monkeypatch.setattr(spreadsheets, "_open_xlsx", lambda _path: book)
        assert _extract_xlsx_authored_at(tmp_path / "book.xlsx") == datetime(
            2020, 1, 1, tzinfo=UTC
        )

    def test_both_candidates_unparseable_yields_none_and_still_closes(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Exhausting the chain returns ``None`` — and closes the handle."""
        book = _FakeWorkbook(created="junk", modified="also junk")
        monkeypatch.setattr(spreadsheets, "_open_xlsx", lambda _path: book)
        assert _extract_xlsx_authored_at(tmp_path / "book.xlsx") is None
        assert book.closed is True


class TestPptxCandidateChain:
    """``created`` outranks ``modified`` for decks too."""

    def test_created_wins_over_modified(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """The deck's ``created`` is preferred when both are usable."""
        props = _FakeProps(
            created=datetime(2024, 3, 15, tzinfo=UTC),
            modified=datetime(2020, 1, 1, tzinfo=UTC),
        )
        monkeypatch.setattr(
            presentations, "_open_pptx_core_properties", lambda _path: props
        )
        assert _extract_pptx_authored_at(tmp_path / "deck.pptx") == datetime(
            2024, 3, 15, tzinfo=UTC
        )

    def test_falls_through_to_modified(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """An absent ``created`` hands off to ``modified``."""
        props = _FakeProps(created=None, modified=datetime(2020, 1, 1, tzinfo=UTC))
        monkeypatch.setattr(
            presentations, "_open_pptx_core_properties", lambda _path: props
        )
        assert _extract_pptx_authored_at(tmp_path / "deck.pptx") == datetime(
            2020, 1, 1, tzinfo=UTC
        )

    def test_both_candidates_unusable_yields_none(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Exhausting the chain returns ``None``, never a filesystem guess."""
        props = _FakeProps(created="junk", modified=None)
        monkeypatch.setattr(
            presentations, "_open_pptx_core_properties", lambda _path: props
        )
        assert _extract_pptx_authored_at(tmp_path / "deck.pptx") is None


class TestDocxCandidateChain:
    """DOCX walks ``created_date`` then ``modified_date``."""

    def test_created_wins_over_modified(self) -> None:
        """``dcterms:created`` outranks ``dcterms:modified``."""
        result = _parse_docx_authored_at(
            {"created_date": "2024-03-15T08:30:00+00:00", "modified_date": "2020-01-01"}
        )
        assert result == datetime(2024, 3, 15, 8, 30, tzinfo=UTC)

    @pytest.mark.parametrize(
        "created",
        [
            pytest.param(None, id="absent"),
            pytest.param("", id="empty"),
            pytest.param("not-a-date", id="unparseable"),
        ],
    )
    def test_falls_through_to_modified(self, created: str | None) -> None:
        """Absent, blank, and unparseable all fall to ``modified_date``."""
        metadata: dict[str, Any] = {"modified_date": "2020-01-01T00:00:00+00:00"}
        if created is not None:
            metadata["created_date"] = created
        result = _parse_docx_authored_at(metadata)
        assert result == datetime(2020, 1, 1, tzinfo=UTC)

    def test_empty_metadata_yields_none(self) -> None:
        """No candidates at all → ``None``."""
        assert _parse_docx_authored_at({}) is None

    def test_both_candidates_unparseable_yields_none(self) -> None:
        """Every candidate raising still returns ``None``, never a raise."""
        assert (
            _parse_docx_authored_at(
                {"created_date": "not-a-date", "modified_date": "also-not-a-date"}
            )
            is None
        )


class TestPdfDateParsing:
    """The ISO 32000 ``D:`` date grammar, and everything it rejects."""

    def test_full_form_with_apostrophe_offset(self) -> None:
        """``+05'00'`` is normalised to ``+0500`` so ``%z`` parses it."""
        result = _parse_pdf_date("D:20240315083000+05'00'")
        assert result is not None
        offset = result.utcoffset()
        assert offset is not None
        assert offset.total_seconds() == 5 * 3600
        assert (result.year, result.month, result.day) == (2024, 3, 15)

    def test_trailing_z_is_treated_as_utc(self) -> None:
        """A ``Z`` suffix means UTC, not a parse failure."""
        assert _parse_pdf_date("D:20240315083000Z") == datetime(
            2024, 3, 15, 8, 30, tzinfo=UTC
        )

    def test_lowercase_z_is_treated_as_utc(self) -> None:
        """Real-world writers emit a lowercase ``z`` too."""
        assert _parse_pdf_date("D:20240315083000z") == datetime(
            2024, 3, 15, 8, 30, tzinfo=UTC
        )

    def test_seconds_precision_without_offset_defaults_to_utc(self) -> None:
        """A naive PDF date is UTC-anchored per FEAT-031, not localised."""
        assert _parse_pdf_date("D:20240315083045") == datetime(
            2024, 3, 15, 8, 30, 45, tzinfo=UTC
        )

    def test_date_only_truncation(self) -> None:
        """``%Y%m%d`` is the shortest accepted form."""
        assert _parse_pdf_date("D:20240315") == datetime(2024, 3, 15, tzinfo=UTC)

    def test_missing_d_prefix_still_parses(self) -> None:
        """The ``D:`` prefix is optional in practice."""
        assert _parse_pdf_date("20240315") == datetime(2024, 3, 15, tzinfo=UTC)

    def test_surrounding_whitespace_is_stripped(self) -> None:
        """Leading/trailing whitespace does not defeat the parse."""
        assert _parse_pdf_date("  D:20240315  ") == datetime(2024, 3, 15, tzinfo=UTC)

    @pytest.mark.parametrize(
        "raw",
        [
            pytest.param("", id="empty"),
            pytest.param("D:", id="prefix-only"),
            pytest.param("not a date", id="prose"),
            pytest.param("D:2024", id="year-only"),
            pytest.param("D:20240230", id="february-30th"),
            pytest.param("D:2024-03-15", id="iso-hyphens"),
        ],
    )
    def test_unparseable_forms_yield_none(self, raw: str) -> None:
        """No candidate format matches → ``None``, never a guess."""
        assert _parse_pdf_date(raw) is None

    @pytest.mark.parametrize(
        ("raw", "expected"),
        [
            pytest.param(
                "D:202403150830",
                datetime(2024, 3, 15, 8, 3, tzinfo=UTC),
                id="twelve-digits-reads-as-8:03:00",
            ),
            pytest.param(
                "D:2024031",
                datetime(2024, 3, 1, tzinfo=UTC),
                id="nine-digits-loses-a-digit",
            ),
            pytest.param(
                "D:20241345",
                datetime(2024, 1, 3, 4, 5, tzinfo=UTC),
                id="impossible-month-reparsed-as-january",
            ),
        ],
    )
    def test_odd_digit_counts_currently_yield_a_wrong_date(
        self, raw: str, expected: datetime
    ) -> None:
        """Pin the loose-``strptime`` hazard as it behaves **today**.

        ``datetime.strptime`` lets each numeric directive match one *or*
        two digits, so ``%Y%m%d%H%M%S`` re-segments a malformed run of
        digits into a plausible-looking but **wrong** date instead of
        rejecting it: ``D:20241345`` becomes 2024-01-03 04:05. That
        contradicts the FEAT-031 "never guess" contract, but tightening
        the grammar risks rejecting real-world PDFs that parse correctly
        today, so it is deliberately **not** changed here — see the
        follow-up issue referenced in the PR body.

        These assertions exist so that a future tightening shows up as a
        deliberate, reviewed change to this table rather than as silent
        drift.
        """
        assert _parse_pdf_date(raw) == expected


class TestPdfCandidateChain:
    """PDF walks ``creationdate`` then ``moddate``."""

    def test_creationdate_wins_over_moddate(self) -> None:
        """The ordered chain prefers ``/CreationDate``."""
        result = _parse_pdf_authored_at(
            {"creationdate": "D:20240315", "moddate": "D:20200101"}
        )
        assert result == datetime(2024, 3, 15, tzinfo=UTC)

    def test_falls_through_to_moddate_when_creationdate_absent(self) -> None:
        """The issue's own worked example: ``moddate`` alone resolves."""
        result = _parse_pdf_authored_at({"moddate": "D:20240102"})
        assert result is not None
        assert result.day == 2
        assert result == datetime(2024, 1, 2, tzinfo=UTC)

    def test_falls_through_when_creationdate_is_unparseable(self) -> None:
        """A junk ``/CreationDate`` does not poison the fallback."""
        result = _parse_pdf_authored_at(
            {"creationdate": "garbage", "moddate": "D:20200101"}
        )
        assert result == datetime(2020, 1, 1, tzinfo=UTC)

    def test_blank_creationdate_is_skipped(self) -> None:
        """An empty string is falsy, so the loop skips before parsing."""
        result = _parse_pdf_authored_at({"creationdate": "", "moddate": "D:20200101"})
        assert result == datetime(2020, 1, 1, tzinfo=UTC)

    def test_both_candidates_unparseable_yields_none(self) -> None:
        """Exhausting the chain returns ``None``."""
        assert (
            _parse_pdf_authored_at({"creationdate": "junk", "moddate": "also junk"})
            is None
        )

    def test_empty_metadata_yields_none(self) -> None:
        """No info dictionary at all → ``None``."""
        assert _parse_pdf_authored_at({}) is None


class TestRtfCandidateChain:
    r"""RTF walks ``\creatim`` then ``\revtim``."""

    def test_creatim_wins_over_revtim(self) -> None:
        r"""``\creatim`` is the authored date; ``\revtim`` is the fallback."""
        raw = (
            rb"{\rtf1{\creatim\yr2024\mo3\dy15\hr8\min30\sec0}{\revtim\yr2020\mo1\dy1}}"
        )
        assert _parse_rtf_authored_at(raw) == datetime(2024, 3, 15, 8, 30, tzinfo=UTC)

    def test_falls_through_to_revtim_when_creatim_incomplete(self) -> None:
        r"""``\creatim`` lacking a day is incomplete, so ``\revtim`` wins."""
        raw = rb"{\rtf1{\creatim\yr2024\mo3}{\revtim\yr2020\mo1\dy1}}"
        assert _parse_rtf_authored_at(raw) == datetime(2020, 1, 1, tzinfo=UTC)

    def test_falls_through_when_creatim_fields_are_impossible(self) -> None:
        r"""An out-of-range ``\mo`` raises in ``datetime``; the arm swallows it."""
        raw = rb"{\rtf1{\creatim\yr2024\mo13\dy45}{\revtim\yr2020\mo1\dy1}}"
        assert _parse_rtf_authored_at(raw) == datetime(2020, 1, 1, tzinfo=UTC)

    def test_negative_year_is_rejected(self) -> None:
        r"""A negative ``\yr`` is out of ``datetime``'s range, not a crash."""
        raw = rb"{\rtf1{\creatim\yr-5\mo1\dy1}}"
        assert _parse_rtf_authored_at(raw) is None

    def test_no_date_control_words_yields_none(self) -> None:
        """A dateless RTF stream returns ``None``."""
        assert _parse_rtf_authored_at(rb"{\rtf1 plain text only}") is None

    def test_both_candidates_incomplete_yields_none(self) -> None:
        """Neither control word carries year+month+day → ``None``."""
        raw = rb"{\rtf1{\creatim\yr2024}{\revtim\mo1\dy1}}"
        assert _parse_rtf_authored_at(raw) is None

    def test_non_ascii_bytes_do_not_crash(self) -> None:
        """The stream is decoded with ``errors="replace"``; junk is tolerated."""
        raw = b"\xff\xfe{\\rtf1{\\creatim\\yr2024\\mo3\\dy15}}"
        assert _parse_rtf_authored_at(raw) == datetime(2024, 3, 15, tzinfo=UTC)

    def test_missing_time_fields_default_to_midnight(self) -> None:
        """Absent hour/minute/second default to zero, not to "now"."""
        raw = rb"{\rtf1{\creatim\yr2024\mo3\dy15}}"
        assert _parse_rtf_authored_at(raw) == datetime(2024, 3, 15, tzinfo=UTC)


class TestHtmlArmInDocumentIngestor:
    """The document ingestor's HTML arm swallows its own parse failure."""

    def test_unparseable_meta_date_yields_none(self, tmp_path: Path) -> None:
        """A present-but-unparseable ``<meta>`` date is not a crash.

        The extraction chain finds the string, ``parse_authored_at``
        rejects it, and the arm returns ``None`` rather than letting the
        ``ValueError`` drop the whole document.
        """
        page = tmp_path / "post.html"
        page.write_text(
            '<html><head><meta name="date" content="sometime last spring">'
            "</head><body><p>Body.</p></body></html>",
            encoding="utf-8",
        )
        ingestor = documents.DocumentIngestor()
        fragments = ingestor.parse(ingestor.discover(page)[0])
        assert len(fragments) == 1
        assert fragments[0].metadata["authored_at"] is None
        assert "Body." in fragments[0].content

    def test_parseable_meta_date_is_used(self, tmp_path: Path) -> None:
        """The positive arm, so the swallow is not masking a total failure."""
        page = tmp_path / "post.html"
        page.write_text(
            '<html><head><meta property="article:published_time" '
            'content="2024-03-15T08:30:00+00:00"></head>'
            "<body><p>Body.</p></body></html>",
            encoding="utf-8",
        )
        ingestor = documents.DocumentIngestor()
        fragments = ingestor.parse(ingestor.discover(page)[0])
        assert fragments[0].metadata["authored_at"] == datetime(
            2024, 3, 15, 8, 30, tzinfo=UTC
        )


class TestIngestorLevelDegradation:
    """End-to-end: a dateless source still ingests, with ``authored_at`` unset."""

    def test_csv_ingests_without_authored_at(self, tmp_path: Path) -> None:
        """CSV has no core properties, so the honest answer is ``None``."""
        csv_path = tmp_path / "rows.csv"
        csv_path.write_text("name,qty\nwidget,3\n", encoding="utf-8")
        ingestor = SpreadsheetIngestor(backend=OpenpyxlBackend())
        fragments = ingestor.parse(ingestor.discover(csv_path)[0])
        assert len(fragments) == 1
        assert fragments[0].metadata["authored_at"] is None

    def test_txt_ingests_without_authored_at(self, tmp_path: Path) -> None:
        """TXT is the canonical "no embedded metadata" case."""
        txt = tmp_path / "note.txt"
        txt.write_text("just some prose\n", encoding="utf-8")
        ingestor = documents.DocumentIngestor()
        fragments = ingestor.parse(ingestor.discover(txt)[0])
        assert fragments[0].metadata["authored_at"] is None

    def test_presentation_ingestor_survives_missing_dep(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A deck whose backend cannot read dates still parses its slides."""
        deck = tmp_path / "deck.pptx"
        deck.write_bytes(b"stub")
        monkeypatch.setattr(
            presentations, "_open_pptx_core_properties", lambda _path: None
        )
        ingestor = PresentationIngestor(backend=_StubDeckBackend())
        fragments = ingestor.parse(ingestor.discover(deck)[0])
        assert len(fragments) == 1
        assert fragments[0].metadata["authored_at"] is None


class _StubDeckBackend:
    """Presentation backend returning one fixed slide, no optional deps."""

    def is_available(self) -> bool:
        """Stubs are always available."""
        return True

    def read_presentation(self, path: Path) -> presentations.PresentationData:
        """Return a single-slide deck named after *path*."""
        return presentations.PresentationData(
            title=path.stem,
            slides=(
                presentations.SlideData(
                    index=1, title="Only slide", body="Body", notes=""
                ),
            ),
        )
